"""
Shared OpenPyXL Styling Constants and Table Helpers for Official Excel Reports (USD & CLP).
Centralizes corporate colors, typography, borders, number formats, and row/header writers.
"""

from typing import Any, Iterable, Sequence
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

NAVY_FILL = PatternFill(start_color="0C0C42", end_color="0C0C42", fill_type="solid")
GOLD_FILL = PatternFill(start_color="996909", end_color="996909", fill_type="solid")
LIGHT_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

WHITE_BOLD = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
NAVY_BOLD = Font(name="Calibri", size=11, bold=True, color="0C0C42")
RED_BOLD = Font(bold=True, color="DC2626")
DIFF_ALERT_FONT = RED_BOLD

THIN_BORDER = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1"),
)

CLP_FORMAT = "$#,##0"
USD_FORMAT = "$#,##0.00"


def write_header_row(
    ws: Worksheet,
    row_idx: int,
    headers: Sequence[str],
    fill: PatternFill = NAVY_FILL,
    center: bool = False,
    wrap_text: bool = False,
) -> None:
    """Writes a styled table header row onto an OpenPyXL worksheet."""
    alignment = (
        Alignment(horizontal="center", vertical="center", wrap_text=wrap_text)
        if (center or wrap_text)
        else None
    )
    for col_idx, header_text in enumerate(headers, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=header_text)
        cell.fill = fill
        cell.font = WHITE_BOLD
        if alignment:
            cell.alignment = alignment


def write_data_row(
    ws: Worksheet,
    row_idx: int,
    values: Sequence[Any],
    money_columns: Sequence[int] = (),
    red_columns: Sequence[int] = (),
    number_format: str = CLP_FORMAT,
) -> None:
    """Writes a bordered data row with optional currency formatting and discrepancy highlighting."""
    for col_idx, value in enumerate(values, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        cell.border = THIN_BORDER
        if col_idx in money_columns:
            cell.number_format = number_format
        if col_idx in red_columns and value:
            cell.font = RED_BOLD


def auto_adjust_columns(worksheets: Iterable[Worksheet], min_width: int = 12, max_width: int = 45) -> None:
    """Auto-adjusts column widths based on the first 40 rows of each worksheet."""
    for ws in worksheets:
        for col in ws.columns:
            col_letter = col[0].column_letter
            max_len = max((len(str(cell.value or "")) for cell in col[:40]), default=0)
            ws.column_dimensions[col_letter].width = max(min_width, min(max_len + 3, max_width))

