"""
Excel Report Exporter for the CLP (Pesos) Card Reconciliation (Phase 2).
Generates the official styled workbook from a committed ReconciliationProcess,
using shared styling helpers from `excel_styles.py`.
"""

import io
from datetime import date
import openpyxl
from openpyxl.styles import Font

from ..models import ReconciliationProcess
from .excel_styles import (
    GOLD_FILL,
    LIGHT_FILL,
    NAVY_BOLD,
    NAVY_FILL,
    auto_adjust_columns,
    write_data_row,
    write_header_row,
)

# Most actionable exceptions first; out-of-period sales are only pending another Bordero.
EXCEPTION_ORDER = {
    "NOT_FOUND_IN_ERP": 0,
    "AMOUNT_MISMATCH": 1,
    "CARD_TYPE_MISMATCH": 2,
    "OUT_OF_PERIOD": 3,
}

MODALITY_LABELS = {
    "CLP_CREDIT": "Crédito",
    "CLP_DEBIT": "Débito",
    "CLP_PREPAID": "Prepago",
    "CREDIT": "Crédito",
    "DEBIT": "Débito",
    "PREPAID": "Prepago",
}


def _split_status_detail(detail: str):
    """system_status_detail is stored as 'Medio · Tipo de error · Acción sugerida'."""
    parts = (detail or "").split(" · ", 2)
    return (parts + ["", "", ""])[:3]


def generate_committed_clp_excel_bytes(process_id: int) -> io.BytesIO:
    """Builds the styled 3-sheet Excel report from a committed CLP ReconciliationProcess."""
    process = ReconciliationProcess.objects.select_related("user").get(id=process_id, process_currency="CLP")
    validations = list(process.bank_validations.select_related("bank_movement").order_by("payment_modality"))
    summaries = list(process.summaries.filter(currency_table="CLP").select_related("card_type"))
    cross_matches = list(
        process.cross_match_results.select_related("card_type", "transbank_transaction").order_by("sale_date", "authorization_code")
    )

    wb = openpyxl.Workbook()

    # Sheet 1: Resumen Pesos
    ws = wb.active
    ws.title = "Resumen Pesos"
    ws.cell(
        row=1,
        column=1,
        value=f"HOTEL PLAZA SAN FRANCISCO — REPORTE OFICIAL CONCILIACIÓN PESOS (PROCESO #{process.id})",
    ).font = Font(size=14, bold=True, color="0C0C42")
    exec_str = process.execution_date.strftime("%d/%m/%Y %H:%M") if process.execution_date else ""
    ws.cell(
        row=2,
        column=1,
        value=f"Periodo: {process.accounting_period or 'Cierre Pesos'} | Auditor: {process.user.username} | Fecha Registro BD: {exec_str}",
    ).font = Font(size=10, italic=True)

    ws.cell(row=4, column=1, value="1. VALIDACIÓN DE ABONOS TRANSBANK EN CARTOLA").font = NAVY_BOLD
    write_header_row(ws, 5, ["Medio de pago", "Abono TBK (CLP)", "Abono Cartola (CLP)", "Diferencia (CLP)", "Fecha Cartola", "Estado"], fill=NAVY_FILL, center=True, wrap_text=True)
    row = 6
    for validation in validations:
        difference = float(validation.transbank_calculated_deposit - (validation.bank_statement_deposit or 0))
        write_data_row(
            ws,
            row,
            [
                MODALITY_LABELS.get(validation.payment_modality, validation.payment_modality),
                float(validation.transbank_calculated_deposit),
                float(validation.bank_statement_deposit or 0),
                difference,
                validation.deposit_date.strftime("%d/%m/%Y") if validation.deposit_date else "No detectado",
                "Confirmado en Cartola" if validation.is_matched else "Pendiente en Cartola",
            ],
            money_columns=(2, 3, 4),
            red_columns=(4,),
        )
        row += 1

    row += 2
    ws.cell(row=row, column=1, value="2. CUADRATURA POR MEDIO DE PAGO Y MARCA (TRANSBANK VS BORDERÓ)").font = NAVY_BOLD
    row += 1
    write_header_row(ws, row, ["Medio · Marca", "Ventas TBK (CLP)", "Ventas Borderó (CLP)", "Diferencia (CLP)", "Comisión + IVA (CLP)", "Abono Neto (CLP)"], fill=NAVY_FILL, center=True, wrap_text=True)
    row += 1
    totals = [0.0] * 5
    for summary in sorted(summaries, key=lambda s: (s.payment_modality or "", s.card_type.name)):
        values = [
            float(summary.transbank_sales_total),
            float(summary.erp_sales_total),
            float(summary.difference),
            float(summary.total_commissions),
            float(summary.net_deposit_total),
        ]
        totals = [t + v for t, v in zip(totals, values)]
        label = f"{MODALITY_LABELS.get(summary.payment_modality, summary.payment_modality)} · {summary.card_type.name}"
        write_data_row(ws, row, [label, *values], money_columns=(2, 3, 4, 5, 6), red_columns=(4,))
        row += 1
    write_data_row(ws, row, ["TOTAL GENERAL CLP", *totals], money_columns=(2, 3, 4, 5, 6))
    for col in range(1, 7):
        ws.cell(row=row, column=col).fill = LIGHT_FILL
        ws.cell(row=row, column=col).font = NAVY_BOLD
    for col_letter in "ABCDEF":
        ws.column_dimensions[col_letter].width = 26

    # Sheet 2: Vouchers Cuadrados
    ws_matched = wb.create_sheet(title="Vouchers Cuadrados")
    write_header_row(ws_matched, 1, [
        "Autorización TBK", "Código en Borderó", "Medio", "Tarjeta", "Fecha Venta", "Huésped", "Hab.",
        "Cajero", "Monto TBK", "Monto Borderó", "Diferencia", "Diagnóstico",
    ], fill=NAVY_FILL, center=True, wrap_text=True)
    row = 2
    for match in (m for m in cross_matches if m.system_result == "MATCHED"):
        modality = match.transbank_transaction.payment_modality if match.transbank_transaction else ""
        write_data_row(
            ws_matched,
            row,
            [
                match.authorization_code or "",
                match.document_number or match.authorization_code or "",
                MODALITY_LABELS.get(modality, modality),
                match.card_type.name if match.card_type else "",
                match.sale_date.strftime("%d/%m/%Y") if match.sale_date else "—",
                match.guest_name or "",
                match.room_number or "",
                match.cashier_username or "",
                float(match.transbank_amount),
                float(match.erp_amount_clp),
                float(match.difference_clp),
                match.system_observation or "",
            ],
            money_columns=(9, 10, 11),
        )
        row += 1

    # Sheet 3: Reporte de Excepciones (sorted by severity, then date)
    ws_exc = wb.create_sheet(title="Reporte de Excepciones")
    write_header_row(ws_exc, 1, [
        "Tipo de error",
        "Fecha de origen",
        "Medio de pago",
        "Código en Transbank",
        "Código en Borderó",
        "Monto Transbank (CLP)",
        "Monto Borderó (CLP)",
        "Monto discrepante (CLP)",
        "Huésped / Hab. / Cajero",
        "Acción sugerida",
        "Detalle",
    ], fill=GOLD_FILL, center=True, wrap_text=True)
    exceptions = sorted(
        (m for m in cross_matches if m.system_result != "MATCHED"),
        key=lambda m: (EXCEPTION_ORDER.get(m.system_result, 9), m.sale_date or date.min, m.authorization_code or ""),
    )
    row = 2
    for match in exceptions:
        modality_label, error_type, action = _split_status_detail(match.system_status_detail)
        context = " · ".join(
            part for part in (match.guest_name, f"Hab. {match.room_number}" if match.room_number and match.room_number != "—" else "", match.cashier_username)
            if part and part != "—"
        )
        write_data_row(
            ws_exc,
            row,
            [
                error_type or match.system_result,
                match.sale_date.strftime("%d/%m/%Y") if match.sale_date else "—",
                modality_label,
                match.authorization_code or "",
                match.document_number or "—",
                float(match.transbank_amount),
                float(match.erp_amount_clp),
                float(match.difference_clp),
                context or "—",
                action,
                match.system_observation or "",
            ],
            money_columns=(6, 7, 8),
            red_columns=(1, 8),
        )
        row += 1

    auto_adjust_columns((ws_matched, ws_exc), min_width=14, max_width=45)

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream
