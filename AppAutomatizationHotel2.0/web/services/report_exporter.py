"""
Excel Report Exporter for Hotel Plaza San Francisco USD Reconciliation (Phase 1).
Generates styled multi-sheet Excel workbooks directly from persisted database records,
using shared styling helpers from `excel_styles.py`.
"""

import io
import json
from typing import Iterable, List, Sequence
import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.worksheet import Worksheet

from ..models import CrossMatchResult, ReconciliationProcess, ReconciliationSummary
from ..templatetags.currency_filters import format_observation
from .excel_styles import (
    CLP_FORMAT,
    DIFF_ALERT_FONT,
    GOLD_FILL,
    LIGHT_FILL,
    NAVY_BOLD,
    NAVY_FILL,
    THIN_BORDER,
    USD_FORMAT,
    auto_adjust_columns,
    write_header_row,
)


def _write_summary_table(
    ws: Worksheet,
    start_row: int,
    title: str,
    headers: Sequence[str],
    summaries: Iterable[ReconciliationSummary],
    total_label: str,
    num_fmt: str,
    decimals: int,
) -> int:
    """Writes a currency summary table (USD or CLP) and returns the last row index written."""
    ws.cell(row=start_row, column=1, value=title).font = NAVY_BOLD
    write_header_row(ws, start_row + 1, headers, fill=NAVY_FILL, center=True)

    curr_row = start_row + 2
    total_col2 = total_col3 = total_diff = 0.0

    for item in summaries:
        val_col2, val_col3, diff_val = (
            float(item.erp_sales_total),
            float(item.transbank_sales_total),
            float(item.difference),
        )
        total_col2 += val_col2
        total_col3 += val_col3
        total_diff += diff_val

        ws.cell(row=curr_row, column=1, value=item.card_type.name).border = THIN_BORDER
        for col_idx, amount in ((2, val_col2), (3, val_col3), (4, diff_val)):
            cell = ws.cell(row=curr_row, column=col_idx, value=amount)
            cell.number_format, cell.border = num_fmt, THIN_BORDER
            if col_idx == 4 and diff_val != 0:
                cell.font = DIFF_ALERT_FONT
        curr_row += 1

    ws.cell(row=curr_row, column=1, value=total_label).font = NAVY_BOLD
    for col_idx, total_val in ((2, total_col2), (3, total_col3), (4, total_diff)):
        rounded = round(total_val, decimals) if decimals > 0 else round(total_val, 0)
        ws.cell(row=curr_row, column=col_idx, value=rounded).number_format = num_fmt
    for col_idx in range(1, 5):
        ws.cell(row=curr_row, column=col_idx).fill = LIGHT_FILL
        ws.cell(row=curr_row, column=col_idx).border = THIN_BORDER
    return curr_row


def _populate_matched_sheet(ws: Worksheet, cross_matches: List[CrossMatchResult]) -> None:
    headers = [
        "Fecha", "Cruce Voucher (TBK = ERP)", "Cruce Documento (ERP = OP)", "Cajero", "Tarjeta",
        "Monto TBK USD", "Monto Operador USD", "Diferencia USD", "Monto ERP CLP", "Diagnóstico",
    ]
    write_header_row(ws, 1, headers, fill=NAVY_FILL)

    for r_idx, v in enumerate((m for m in cross_matches if m.system_result == "MATCHED"), 2):
        ws.cell(row=r_idx, column=1, value=v.sale_date.strftime("%d/%m/%Y") if v.sale_date else "—")
        ws.cell(row=r_idx, column=2, value=f"TBK: {v.authorization_code} = ERP: {v.authorization_code}" if v.authorization_code else "")
        ws.cell(row=r_idx, column=3, value=f"ERP: {v.document_number} = OP: {v.document_number}" if v.document_number else "")
        ws.cell(row=r_idx, column=4, value=v.cashier_username or "")
        ws.cell(row=r_idx, column=5, value=v.card_type.name if v.card_type else "")
        ws.cell(row=r_idx, column=6, value=float(v.transbank_amount)).number_format = USD_FORMAT
        ws.cell(row=r_idx, column=7, value=float(v.operator_amount_usd)).number_format = USD_FORMAT
        ws.cell(row=r_idx, column=8, value=float(v.difference_usd)).number_format = USD_FORMAT
        ws.cell(row=r_idx, column=9, value=float(v.erp_amount_clp)).number_format = CLP_FORMAT
        ws.cell(row=r_idx, column=10, value=format_observation(v.system_observation or ""))


def _populate_discrepancies_sheet(ws: Worksheet, cross_matches: List[CrossMatchResult]) -> None:
    headers = [
        "Autorización", "Documento", "Huésped", "Habitación", "Tarjeta",
        "Monto TBK USD", "Monto Operador USD", "Diferencia USD", "Diagnóstico",
    ]
    write_header_row(ws, 1, headers, fill=GOLD_FILL)

    for r_idx, v in enumerate((m for m in cross_matches if m.system_result != "MATCHED"), 2):
        ws.cell(row=r_idx, column=1, value=v.authorization_code or "")
        ws.cell(row=r_idx, column=2, value=v.document_number or "")
        ws.cell(row=r_idx, column=3, value=v.guest_name or "")
        ws.cell(row=r_idx, column=4, value=v.room_number or "")
        ws.cell(row=r_idx, column=5, value=v.card_type.name if v.card_type else "")
        ws.cell(row=r_idx, column=6, value=float(v.transbank_amount)).number_format = USD_FORMAT
        ws.cell(row=r_idx, column=7, value=float(v.operator_amount_usd)).number_format = USD_FORMAT
        ws.cell(row=r_idx, column=8, value=float(v.difference_usd)).number_format = USD_FORMAT
        ws.cell(row=r_idx, column=9, value=format_observation(v.system_observation or ""))


def _populate_suggestions_sheet(ws: Worksheet, cross_matches: List[CrossMatchResult]) -> None:
    headers = [
        "#", "Voucher TBK", "Tipo Discrepancia", "Tarjeta", "Monto TBK (US$)", "Fecha POS",
        "Cód. ERP Sugerido", "Doc. ERP", "Monto ERP (CLP)", "Cajero ERP", "Huésped / Habitación",
        "Doc. Operadora", "Monto Op. (US$)", "Saldo Op. (CLP)", "Confianza",
        "Diagnóstico Sugerido", "Motivos Coincidencia", "Nota de Auditoría",
    ]
    write_header_row(ws, 1, headers, fill=NAVY_FILL, center=True)

    center_align = Alignment(horizontal="center")
    sug_count, curr_sug_row = 0, 2

    for v in cross_matches:
        if not v.auditor_comment:
            continue
        try:
            sug = json.loads(v.auditor_comment)
            if not isinstance(sug, dict):
                continue
        except Exception:
            continue

        sug_count += 1
        case_type = "Con Diferencia de Monto" if v.system_result == "AMOUNT_MISMATCH" else "En TBK pero NO en ERP"
        reasons_str = ", ".join(r.get("label", "") for r in sug.get("reasons", []) if isinstance(r, dict))
        confidence = sug.get("confidence", "MEDIA")

        rows_comp = sug.get("comparison_rows", [])
        r_auth = rows_comp[0] if len(rows_comp) > 0 else {}
        r_doc = rows_comp[1] if len(rows_comp) > 1 else {}
        r_clp = rows_comp[3] if len(rows_comp) > 3 else {}

        guest_str = f"{v.guest_name} (Hab. {v.room_number})" if v.guest_name and v.guest_name != "—" else "—"
        sale_date_str = v.sale_date.strftime("%d/%m/%Y") if v.sale_date else "—"

        ws.cell(row=curr_sug_row, column=1, value=sug_count).alignment = center_align
        ws.cell(row=curr_sug_row, column=2, value=v.authorization_code or "").alignment = center_align
        ws.cell(row=curr_sug_row, column=3, value=case_type)
        ws.cell(row=curr_sug_row, column=4, value=v.card_type.name if v.card_type else "")
        ws.cell(row=curr_sug_row, column=5, value=float(v.transbank_amount)).number_format = USD_FORMAT
        ws.cell(row=curr_sug_row, column=6, value=sale_date_str).alignment = center_align
        ws.cell(row=curr_sug_row, column=7, value=r_auth.get("erp_value", "—")).alignment = center_align
        ws.cell(row=curr_sug_row, column=8, value=r_doc.get("erp_value", "—"))
        ws.cell(row=curr_sug_row, column=9, value=float(v.erp_amount_clp)).number_format = CLP_FORMAT
        ws.cell(row=curr_sug_row, column=10, value=sug.get("erp_cashier", v.cashier_username or "—"))
        ws.cell(row=curr_sug_row, column=11, value=guest_str)
        ws.cell(row=curr_sug_row, column=12, value=r_doc.get("op_value", "—"))
        ws.cell(row=curr_sug_row, column=13, value=float(v.operator_amount_usd)).number_format = USD_FORMAT
        ws.cell(row=curr_sug_row, column=14, value=r_clp.get("op_value", ""))

        c_conf = ws.cell(row=curr_sug_row, column=15, value=confidence)
        c_conf.alignment = center_align
        c_conf.font = Font(bold=True, color="059669" if confidence == "ALTA" else "D97706")

        ws.cell(row=curr_sug_row, column=16, value=sug.get("headline", ""))
        ws.cell(row=curr_sug_row, column=17, value=reasons_str)
        ws.cell(row=curr_sug_row, column=18, value=sug.get("manual_check_note", ""))

        for col_i in range(1, 19):
            ws.cell(row=curr_sug_row, column=col_i).border = THIN_BORDER
        curr_sug_row += 1

    if sug_count == 0:
        c_empty = ws.cell(
            row=2,
            column=1,
            value="No se detectaron sugerencias asistidas secundarias para los vouchers con diferencia o no ubicados.",
        )
        c_empty.font = Font(italic=True, color="64748B")


def generate_committed_usd_excel_bytes(process_id: int) -> io.BytesIO:
    """Builds a styled 4-sheet openpyxl Excel report from a committed ReconciliationProcess."""
    process = ReconciliationProcess.objects.select_related("user").get(id=process_id)
    usd_summaries = list(process.summaries.filter(currency_table="USD").select_related("card_type"))
    clp_summaries = list(process.summaries.filter(currency_table="CLP").select_related("card_type"))
    bank_val = process.bank_validations.first()
    cross_matches = list(process.cross_match_results.select_related("card_type"))

    wb = openpyxl.Workbook()
    ws_summary = wb.active
    ws_summary.title = "Resumen Dólares"

    ws_summary.cell(
        row=1,
        column=1,
        value=f"HOTEL PLAZA SAN FRANCISCO — REPORTE OFICIAL CONCILIACIÓN DÓLARES (PROCESO #{process.id})",
    ).font = Font(size=14, bold=True, color="0C0C42")

    exec_str = process.execution_date.strftime("%d/%m/%Y %H:%M") if process.execution_date else ""
    bank_status = "Confirmado en Cartola" if (bank_val and bank_val.is_matched) else "Pendiente en Cartola"
    deposit_usd = float(bank_val.transbank_calculated_deposit) if bank_val else 0.0
    ws_summary.cell(
        row=2,
        column=1,
        value=(
            f"Periodo: {process.accounting_period or 'Cierre USD'} | "
            f"Auditor: {process.user.username} | "
            f"Fecha Registro BD: {exec_str} | "
            f"Abono TBK: US$ {deposit_usd:,.2f} ({bank_status})"
        ),
    ).font = Font(size=10, italic=True)

    last_usd_row = _write_summary_table(
        ws_summary,
        4,
        "1. CUADRATURA POR OPERADOR EN DÓLARES (US$)",
        ["Operador", "Total TBK (US$)", "Total Operador (US$)", "Diferencia (US$)"],
        usd_summaries,
        "TOTAL GENERAL USD",
        USD_FORMAT,
        2,
    )
    _write_summary_table(
        ws_summary,
        last_usd_row + 3,
        "2. CUADRATURA EN MONEDA NACIONAL (PESOS CLP)",
        ["Operador", "Total ERP (CLP)", "Total Operador (CLP)", "Diferencia (CLP)"],
        clp_summaries,
        "TOTAL GENERAL CLP",
        CLP_FORMAT,
        0,
    )
    for col_letter in ("A", "B", "C", "D"):
        ws_summary.column_dimensions[col_letter].width = 28

    ws_matched = wb.create_sheet(title="Vouchers Ubicados")
    _populate_matched_sheet(ws_matched, cross_matches)

    ws_disc = wb.create_sheet(title="Discrepancias y No Ubicados")
    _populate_discrepancies_sheet(ws_disc, cross_matches)

    ws_sug = wb.create_sheet(title="Sugerencias Asistidas")
    _populate_suggestions_sheet(ws_sug, cross_matches)

    auto_adjust_columns((ws_matched, ws_disc, ws_sug))

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream
