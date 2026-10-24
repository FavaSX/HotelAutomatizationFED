import io
from datetime import date
from decimal import Decimal

import openpyxl
from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse

from .models import (
    BankStatementMovement,
    CrossMatchResult,
    ErpTransaction,
    ReconciliationProcess,
    SystemConfiguration,
    TransbankTransaction,
)
from .services import calculate_clp_preview, commit_clp_preview_to_database, detect_and_align_clp_files
from .services.clp_reconciliation import get_clp_tolerance
from .services.column_detector import normalize_auth_code, safe_date
from .services.clp_report_exporter import generate_committed_clp_excel_bytes


# ------------------------------------------------------------------------------
# Synthetic Excel files with the same layout as the real hotel reports
# ------------------------------------------------------------------------------

def _xlsx(name: str, rows: dict) -> io.BytesIO:
    """rows = {excel_row: {column_letter: value}}"""
    wb = openpyxl.Workbook()
    ws = wb.active
    for row_number, cells in rows.items():
        for column, value in cells.items():
            ws[f"{column}{row_number}"] = value
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    stream.name = name
    return stream


def build_tbk_settlement(name: str, medium: str, sales: list) -> io.BytesIO:
    """
    Transbank 'Abonos por día'. sales = [(fecha, local, monto, tipo, cuota, valido, comision, iva, tarjeta, codigo)].
    """
    total_sales = sum(s[5] for s in sales)
    commission = sum(s[6] for s in sales)
    vat = sum(s[7] for s in sales)
    rows = {
        5: {"B": "Reporte:", "C": "Abonos por día", "E": "Periodo de consulta:", "F": "08 septiembre 2026"},
        9: {"B": "Empresa:", "C": "RUT:", "D": "Cuenta corriente:", "E": "Moneda:", "F": "Medio de pago:"},
        10: {"B": "HOTELERA SAN FRANCISCO S.A.", "C": "99.511.100-4", "D": "BANCO DE CHILE 6211002", "E": "Pesos", "F": medium},
        17: {"B": "Total ventas (+)", "C": total_sales},
        21: {"B": "Comisión Transbank (-)", "C": commission},
        22: {"B": "IVA de comisión Transbank (-)", "C": vat},
        26: {"B": "Total abono", "C": total_sales - commission - vat},
        29: {"B": "Fecha de venta", "C": "Tipo de movimiento", "D": "Código de comercio", "E": "Nombre local", "F": "Medio de pago", "G": "Ventas", "Y": "Identificador"},
        30: {
            "G": "Monto original de venta (+)", "I": "Tipo de venta/cuota", "J": "N° de cuota",
            "K": "Monto venta válido para abono (+)", "L": "Comisión Transbank (-)", "M": "IVA Comisión Transbank (-)",
            "N": "Total abono (=)", "Y": "N° de tarjeta", "Z": "Código de autorización de venta",
        },
    }
    for offset, (fecha, local, monto, tipo, cuota, valido, comision, iva, tarjeta, codigo) in enumerate(sales):
        rows[31 + offset] = {
            "B": fecha, "C": "Venta", "D": "27941095", "E": f"{local}, AVENIDA LIBERTADOR BERNARDO O'HIGGINS 816",
            "F": medium, "G": monto, "I": tipo, "J": cuota, "K": valido, "L": comision, "M": iva,
            "N": valido - comision - iva, "Y": tarjeta, "Z": codigo,
        }
    return _xlsx(name, rows)


def build_bordero(payments: list) -> io.BytesIO:
    """Bordero rows: (hab, cuenta, cod, descripcion, factura, valor, documento, fecha, hora, usuario, huesped)."""
    header = {"A": "Hab", "B": "Tipo Hab", "C": "Reserva", "D": "Cuenta", "E": "Cod. Deb.", "F": "Descripción",
              "H": "Factura", "I": "Valor", "J": "Documento", "K": "Fecha", "L": "Hora", "M": "Usuario", "N": "Designación"}
    rows = {
        1: {"A": "HOTELERA SAN FRANCISCO S.A"},
        2: {"A": "Cargos de Débitos y Créditos"},
        4: {"A": "Periodo de:", "B": "04/09/2026", "C": "hasta", "D": "07/09/2026"},
        5: header,
        6: {"A": 1002, "D": 1785942, "E": "CNBQ", "F": "1/2 Pension Alimentos", "I": 20900, "J": "UH: 1002", "K": "07/09/2026"},
        7: {"A": "Total VHF:", "B": 20900},
    }
    for offset, (hab, cuenta, cod, desc, factura, valor, documento, fecha, hora, usuario, huesped) in enumerate(payments):
        rows[8 + offset] = {"A": hab, "D": cuenta, "E": cod, "F": desc, "G": factura, "I": valor,
                            "J": documento, "K": fecha, "L": hora, "M": usuario, "N": huesped}
    return _xlsx("BORDERO 04-09 A 07-09.xlsx", rows)


def build_cartola(deposits: list) -> io.BytesIO:
    rows = {
        11: {"B": "Cuenta N°:", "C": "00-000-62110-02"},
        12: {"B": "Moneda:", "C": "Pesos Chilenos (CLP)"},
        15: {"B": "Saldo Disponible", "C": "Saldo Contable"},
        22: {"B": "Fecha", "D": "Descripción", "F": "Canal o Sucursal", "G": "Nro. Docto.", "H": "Cargos (CLP)", "I": "Abonos (CLP)", "J": "Saldo (CLP)"},
    }
    for offset, (fecha, descripcion, abono) in enumerate(deposits):
        rows[23 + offset] = {"B": fecha, "D": descripcion, "F": "Oficina Central", "I": abono, "J": 1000000}
    return _xlsx("cartola.xlsx", rows)


CREDIT_SALES = [
    # MATCHED: Mastercard credit 405544 (the study guide example)
    ("04 septiembre 2026", "HOTEL PLAZA SAN FRANCISCO", 187172, "Venta $", "S/C", 187172, 4997, 949, "5525 61** **** 5410", "405544"),
    # OUT_OF_PERIOD: sale before the Bordero period
    ("03 septiembre 2026", "RESTAURANT BRISTOL", 8690, "Venta $", "S/C", 8690, 174, 33, "5532 53** **** 4291", "070981"),
    # OUT_OF_PERIOD: installment 3/3 of a July sale (compared with the original amount)
    ("04 julio 2026", "HOTEL PLAZA SAN FRANCISCO", 1354387, "C3C", "3/3", 451463, 11377, 2162, "4260 40** **** 7444", "626135"),
]
DEBIT_SALES = [
    # MATCHED with Excel-dropped leading zeros (TBK 003098 -> Bordero 3098) and a 1 peso difference
    ("07 septiembre 2026", "RESTAURANT BRISTOL", 57420, "Normal", "S/C", 57420, 568, 108, "4547 93** **** 1111", "003098"),
    # AMOUNT_MISMATCH: 1.504 above the ±1.000 maximum
    ("05 septiembre 2026", "RESTAURANT BRISTOL", 115204, "Normal", "S/C", 115204, 1140, 217, "4547 93** **** 2222", "007626"),
    # NOT_FOUND_IN_ERP with typo suggestion: cashier typed 0FOFUAZ
    ("06 septiembre 2026", "HOTEL PLAZA SAN FRANCISCO", 106999, "Normal", "S/C", 106999, 856, 163, "4547 93** **** 8430", "F0FUAZ"),
    # MATCHED within tolerance (704 pesos)
    ("05 septiembre 2026", "HOTEL PLAZA SAN FRANCISCO", 115204, "Normal", "S/C", 115204, 1140, 217, "4547 93** **** 3333", "112233"),
]
PREPAID_SALES = [
    # MATCHED by netting a reversal: -11.000, +11.000, -9.800
    ("05 septiembre 2026", "RESTAURANT BRISTOL", 9800, "-", "S/C", 9800, 97, 18, "5303 12** **** 4444", "823431"),
]
BORDERO_PAYMENTS = [
    (822, 1786484, "MC$", "Master Card $", 53895, -187172, 405544, "04/09/2026", "03:11:09", "CDAZA", "RODRIGO ISLA"),
    (None, 1786798, "RCO", "Red Compra", 259900, -57421, 3098, "07/09/2026", "13:00:00", "BFREDERIC", "PARC"),
    (None, 1786853, "RCO", "Red Compra", 259924, -113700, 7626, "05/09/2026", "15:38:37", "SCORREA", "PARC"),
    (825, 1787015, "RCO", "Red Compra", 259977, -106999, "0FOFUAZ", "06/09/2026", "19:36:48", "MSALAS", "KATHARINA WINCO"),
    (None, 1786900, "RCO", "Red Compra", 259930, -114500, 112233, "05/09/2026", "16:00:00", "SCORREA", "PARC"),
    (None, 1786794, "RCO", "Red Compra", None, 11000, 823431, "05/09/2026", "09:26:26", "SCORREA", "PARC"),
    (None, 1786794, "RCO", "Red Compra", None, -11000, 823431, "05/09/2026", "09:25:33", "SCORREA", "PARC"),
    (None, 1786793, "RCO", "Red Compra", 259889, -9800, 823431, "05/09/2026", "09:23:22", "SCORREA", "PARC"),
    # A Phase 1 dollar payment and cash must be ignored
    (602, 1715625, "VUS$", "Visa US$", 99965, -107984, 112194, "05/09/2026", "19:28:02", "TMENDOZA", "MARK WILLIS"),
    (None, 1786000, "E$", "Efectivo $", None, -50000, None, "05/09/2026", "10:00:00", "SCORREA", "PARC"),
]


def _deposit(sales):
    return sum(s[5] - s[6] - s[7] for s in sales)


def build_clp_files() -> dict:
    return {
        "erp_file": build_bordero(BORDERO_PAYMENTS),
        "bank_statement_file": build_cartola([
            ("08/09/2026", "Pago: Abonos Realizados Por Transbank Al Comercio Afilia 0966893109", _deposit(PREPAID_SALES)),
            ("08/09/2026", "Pago: Abonos Realizados Por Transbank Al Comercio Afilia 0966893109", _deposit(DEBIT_SALES)),
            ("08/09/2026", "Pago: Proveedores 0856339009", 580075),
            ("08/09/2026", "Pago: Abono Comer. Tc Tbk 0966893109", _deposit(CREDIT_SALES)),
        ]),
        "transbank_credit_file": build_tbk_settlement("CREDITOS TBK 08-09.xlsx", "Crédito", CREDIT_SALES),
        "transbank_debit_file": build_tbk_settlement("DEBITOS TBK 08-09.xlsx", "Débito", DEBIT_SALES),
        "transbank_prepaid_file": build_tbk_settlement("PREPAGO TBK 08-09.xlsx", "Prepago", PREPAID_SALES),
    }


def _by_code(preview, bucket):
    return {v["authorization_code"]: v for v in preview[bucket]}


# ------------------------------------------------------------------------------
# Tests
# ------------------------------------------------------------------------------

class ClpHelpersTests(TestCase):
    def test_normalize_auth_code_drops_leading_zeros(self):
        self.assertEqual(normalize_auth_code("003098"), normalize_auth_code(3098))
        self.assertEqual(normalize_auth_code("0WBNIJW"), "WBNIJW")
        self.assertEqual(normalize_auth_code(405544.0), "405544")

    def test_safe_date_reads_spanish_transbank_dates(self):
        self.assertEqual(safe_date("04 septiembre 2026"), date(2026, 9, 4))
        self.assertEqual(safe_date("04/09/2026"), date(2026, 9, 4))

    def test_tolerance_comes_from_system_configuration(self):
        config = SystemConfiguration.objects.first()
        self.assertEqual(get_clp_tolerance(), Decimal("1000"))
        config.clp_rounding_tolerance = Decimal("700")
        config.save()
        self.assertEqual(get_clp_tolerance(), Decimal("700"))


class ClpPreviewTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.preview = calculate_clp_preview(**build_clp_files())

    def test_bank_deposits_match_each_settlement(self):
        validations = {v["payment_modality"]: v for v in self.preview["bank_validations"]}
        self.assertTrue(all(v["is_matched"] for v in validations.values()))
        self.assertIn("TC TBK", validations["CREDIT"]["bank_description"].upper())
        self.assertEqual(validations["DEBIT"]["bank_deposit_amount"], _deposit(DEBIT_SALES))
        self.assertEqual(validations["PREPAID"]["bank_deposit_amount"], _deposit(PREPAID_SALES))

    def test_traced_voucher_405544_matches_bordero_row(self):
        voucher = _by_code(self.preview, "matched_vouchers")["405544"]
        self.assertEqual(voucher["erp_payment_code"], "MC$")
        self.assertEqual(voucher["difference"], 0)
        self.assertEqual(voucher["guest_name"], "RODRIGO ISLA")
        self.assertEqual(voucher["card_code"], "MC")

    def test_leading_zeros_reversals_and_tolerance(self):
        matched = _by_code(self.preview, "matched_vouchers")
        self.assertEqual(matched["003098"]["difference"], -1)
        self.assertEqual(len(matched["823431"]["erp_indices"]), 3)
        self.assertEqual(matched["823431"]["erp_amount"], 9800)
        self.assertEqual(matched["112233"]["difference"], 704)

    def test_amount_above_tolerance_is_a_discrepancy(self):
        voucher = _by_code(self.preview, "discrepancy_vouchers")["007626"]
        self.assertEqual(voucher["system_result"], "AMOUNT_MISMATCH")
        self.assertEqual(voucher["difference"], 1504)

    def test_out_of_period_and_later_installments(self):
        out_of_period = _by_code(self.preview, "out_of_period_vouchers")
        self.assertEqual(set(out_of_period), {"070981", "626135"})
        self.assertIn("Cuota 3/3", out_of_period["626135"]["observation"])

    def test_orphan_typo_gets_high_confidence_suggestion(self):
        voucher = _by_code(self.preview, "orphan_vouchers")["F0FUAZ"]
        self.assertEqual(voucher["suggestion"]["confidence"], "ALTA")
        self.assertEqual(voucher["suggestion"]["erp_row"]["authorization_code"], "0FOFUAZ")
        self.assertEqual(voucher["error_type"], "Error de tipeo")

    def test_only_peso_card_payments_are_read_from_bordero(self):
        codes = {row["payment_code"] for row in self.preview["erp_transactions"]}
        self.assertEqual(codes, {"MC$", "RCO"})


class ClpPersistenceTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user("auditor", password="x")

    def test_commit_export_and_reupload_without_duplicates(self):
        process = commit_clp_preview_to_database(self.user, calculate_clp_preview(**build_clp_files()))
        self.assertEqual(process.process_currency, "CLP")
        self.assertEqual(process.tolerance_applied, Decimal("1000"))
        self.assertEqual(process.bank_validations.filter(is_matched=True).count(), 3)
        self.assertEqual(process.cross_match_results.count(), len(CREDIT_SALES + DEBIT_SALES + PREPAID_SALES))
        counts = (TransbankTransaction.objects.count(), ErpTransaction.objects.count(),
                  CrossMatchResult.objects.count(), BankStatementMovement.objects.count())

        orphan = process.cross_match_results.get(authorization_code="F0FUAZ")
        self.assertEqual(orphan.document_number, "0FOFUAZ")

        workbook = openpyxl.load_workbook(generate_committed_clp_excel_bytes(process.id))
        self.assertEqual(workbook.sheetnames, ["Resumen Pesos", "Vouchers Cuadrados", "Reporte de Excepciones"])
        first_exception = [c.value for c in workbook["Reporte de Excepciones"][2]]
        self.assertEqual(first_exception[0], "Error de tipeo")

        again = commit_clp_preview_to_database(self.user, calculate_clp_preview(**build_clp_files()))
        self.assertEqual(again.id, process.id)
        self.assertEqual(ReconciliationProcess.objects.count(), 1)
        self.assertEqual(
            counts,
            (TransbankTransaction.objects.count(), ErpTransaction.objects.count(),
             CrossMatchResult.objects.count(), BankStatementMovement.objects.count()),
        )


class ToleranceSettingsTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_superuser("admin", password="x")
        self.auditor = User.objects.create_user("auditor", password="x")

    def _post(self, **overrides):
        data = {"usd_rounding_tolerance": "0.50", "clp_rounding_tolerance": "700", **overrides}
        return self.client.post(reverse("settings"), data)

    def test_only_superuser_can_access(self):
        self.client.force_login(self.auditor)
        self.assertRedirects(self.client.get(reverse("settings")), reverse("dashboard"))
        self._post()
        self.assertEqual(SystemConfiguration.get_solo().clp_rounding_tolerance, Decimal("1000"))

    def test_superuser_updates_tolerances(self):
        self.client.force_login(self.admin)
        self.assertEqual(self.client.get(reverse("settings")).status_code, 200)
        self.assertRedirects(self._post(usd_rounding_tolerance="1.25"), reverse("settings"))
        config = SystemConfiguration.get_solo()
        self.assertEqual(config.usd_rounding_tolerance, Decimal("1.25"))
        self.assertEqual(config.clp_rounding_tolerance, Decimal("700"))
        self.assertEqual(config.updated_by, self.admin)
        self.assertEqual(SystemConfiguration.objects.count(), 1)

    def test_rejects_out_of_range_values(self):
        self.client.force_login(self.admin)
        for overrides in (
            {"usd_rounding_tolerance": "5.01"},
            {"clp_rounding_tolerance": "5001"},
            {"clp_rounding_tolerance": "-1"},
        ):
            self.assertEqual(self._post(**overrides).status_code, 200)
        config = SystemConfiguration.get_solo()
        self.assertEqual(config.usd_rounding_tolerance, Decimal("0.50"))
        self.assertEqual(config.clp_rounding_tolerance, Decimal("1000"))


class ClpViewsTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user("auditor", password="x")
        self.client.force_login(self.user)

    def test_detect_files_aligns_the_five_roles(self):
        files = build_clp_files()
        report = detect_and_align_clp_files(list(files.values()))
        self.assertTrue(report["all_matched"])
        roles = {r["role_key"]: r["file_name"] for r in report["aligned_roles"]}
        self.assertEqual(roles["tbk_debit"], "DEBITOS TBK 08-09.xlsx")
        self.assertEqual(roles["erp"], "BORDERO 04-09 A 07-09.xlsx")

    def test_upload_preview_commit_and_export(self):
        files = build_clp_files()
        payload = {
            "clp_erp_file": files["erp_file"],
            "clp_bank_file": files["bank_statement_file"],
            "clp_tbk_credit_file": files["transbank_credit_file"],
            "clp_tbk_debit_file": files["transbank_debit_file"],
            "clp_tbk_prepaid_file": files["transbank_prepaid_file"],
        }
        response = self.client.post(reverse("clp_upload"), payload)
        self.assertRedirects(response, reverse("clp_preview"))

        response = self.client.get(reverse("clp_preview"))
        self.assertContains(response, "F0FUAZ")
        self.assertContains(response, "Ubicadas Exactas")
        self.assertContains(response, "Tolerancia ±$1.000 CLP")
        self.assertNotContains(response, "Borderó histórico (")

        response = self.client.post(reverse("clp_commit"), HTTP_X_REQUESTED_WITH="XMLHttpRequest")
        data = response.json()
        self.assertTrue(data["success"])
        self.assertEqual(self.client.get(data["export_url"]).status_code, 200)

    def test_usd_upload_and_preview_render(self):
        self.assertEqual(self.client.get(reverse("usd_upload")).status_code, 200)

        # Verify USD preview renders the USD tolerance badge when session is populated
        session = self.client.session
        session["usd_reconciliation_preview"] = {
            "accounting_period": "Cierre Dólares — 08/09/2026",
            "tolerance": 0.50,
            "bank_validation": {},
            "counts": {"matched": 0, "discrepancies": 0, "orphans": 0, "total": 0},
            "suggestion_counts": {},
        }
        session.save()
        response = self.client.get(reverse("usd_preview"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Tolerancia ±US$ 0,50 USD")


class DashboardViewTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="contador", password="clave-de-prueba-123")

    def test_requires_login(self):
        response = self.client.get(reverse("dashboard"))
        self.assertRedirects(response, f"{reverse('login')}?next=/")

    def test_empty_state_without_processes(self):
        self.client.force_login(self.user)
        response = self.client.get(reverse("dashboard"))
        self.assertEqual(response.status_code, 200)
        self.assertFalse(response.context["usd"]["has_data"])
        self.assertFalse(response.context["clp"]["has_data"])
        self.assertContains(response, "Todavía no hay cierres en dólares")
        self.assertContains(response, "Todavía no hay cierres en pesos")

    def _clp_closing(self):
        """A saved CLP closing shaped like the ones commit_clp_preview_to_database stores."""
        from .models import BankDepositValidation, CardType, ReconciliationStatus, ReconciliationSummary

        visa, _ = CardType.objects.get_or_create(code="VI", defaults={"name": "Visa"})
        pending, _ = ReconciliationStatus.objects.get_or_create(code="PENDING", defaults={"display_name": "Pendiente"})
        accepted, _ = ReconciliationStatus.objects.get_or_create(code="ACCEPTED_ROUNDING", defaults={"display_name": "Aceptado"})
        process = ReconciliationProcess.objects.create(
            user=self.user, process_currency="CLP", accounting_period="Cierre Pesos — Abono 08/09/2026",
        )
        rows = [
            # (modality, result, amount, status detail)
            ("CREDIT", "MATCHED", 10000, "Crédito · Cuadrado · Sin acción"),
            ("DEBIT", "MATCHED", 20000, "Débito · Cuadrado · Sin acción"),
            ("DEBIT", "NOT_FOUND_IN_ERP", 250000, "Débito · Falta en ERP · Buscar la venta con el cajero y registrarla en el ERP"),
            ("CREDIT", "OUT_OF_PERIOD", 5000, "Crédito · Fuera de período · Buscar en el Borderó del período de la venta"),
        ]
        for i, (modality, result, amount, detail) in enumerate(rows):
            tbk = TransbankTransaction.objects.create(
                process=process, card_type=visa, payment_modality=modality,
                authorization_code=f"A{i}", original_sale_amount=amount,
            )
            CrossMatchResult.objects.create(
                process=process, card_type=visa, transbank_transaction=tbk, authorization_code=f"A{i}",
                transbank_amount=amount, system_result=result, system_status_detail=detail,
                management_status=accepted if result == "MATCHED" else pending,
            )
        for modality, matched in (("CLP_CREDIT", True), ("CLP_DEBIT", True), ("CLP_PREPAID", False)):
            BankDepositValidation.objects.create(
                process=process, payment_modality=modality, transbank_calculated_deposit=1000, is_matched=matched,
            )
        ReconciliationSummary.objects.create(
            process=process, card_type=visa, payment_modality="CLP_DEBIT", currency_table="CLP", total_commissions=350,
        )
        return process

    def test_clp_panel_figures_come_from_saved_closing(self):
        self._clp_closing()
        self.client.force_login(self.user)
        response = self.client.get(reverse("dashboard"), {"moneda": "clp"})

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["active_currency"], "clp")
        clp = response.context["clp"]
        self.assertTrue(clp["has_data"])
        self.assertEqual(clp["last_process"].match_rate, 50.0)
        self.assertEqual(clp["last_counts"]["orphans"], 1)
        self.assertEqual(clp["last_counts"]["out_of_period"], 1)
        self.assertEqual(clp["pending_count"], 2)
        self.assertEqual(clp["pending_amount"], 255000)
        self.assertEqual(clp["out_of_period_amount"], 5000)
        self.assertEqual(clp["last_process"].banks_matched, 2)
        errors = {g["label"]: g for g in clp["pending_by_error"]}
        self.assertEqual(errors["Falta en ERP"]["count"], 1)
        self.assertEqual(errors["Falta en ERP"]["action"], "Buscar la venta con el cajero y registrarla en el ERP")
        modalities = {m["code"]: m for m in clp["modalities"]}
        self.assertEqual((modalities["DEBIT"]["vouchers"], modalities["DEBIT"]["matched"]), (2, 1))
        self.assertEqual(modalities["DEBIT"]["commissions"], 350)
        self.assertEqual(modalities["PREPAID"]["vouchers"], 0)
        self.assertFalse(response.context["usd"]["has_data"])

    def test_defaults_to_clp_panel_when_only_pesos_exist(self):
        self._clp_closing()
        self.client.force_login(self.user)
        response = self.client.get(reverse("dashboard"))
        self.assertEqual(response.context["active_currency"], "clp")


class UsdPreviewViewTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="contador", password="clave-de-prueba-123")
        self.client.force_login(self.user)

    def _store_preview(self, preview):
        session = self.client.session
        session["usd_reconciliation_preview"] = preview
        session.save()

    def test_redirects_to_upload_without_preview(self):
        response = self.client.get(reverse("usd_preview"))
        self.assertRedirects(response, reverse("usd_upload"))

    def test_merges_usd_and_clp_rows_per_franchise(self):
        voucher = {
            "authorization_code": "A10001", "document_number": "9000001", "card_name": "Visa",
            "guest_name": "HUESPED DE PRUEBA", "room_number": "101", "cashier_username": "CAJERO1",
            "gross_amount_usd": 100.0, "operator_amount_usd": 100.11, "difference_usd": -0.11,
            "erp_amount_clp": 95000.0, "observation": "Coincide con SALDO 95000",
        }
        self._store_preview({
            "accounting_period": "Cierre Dólares — Abono de prueba",
            "bank_validation": {"calculated_deposit": 100.0, "bank_deposit_amount": 100.0, "is_matched": True},
            "usd_table": [
                {"card_code": "VI", "card_name": "Visa US$", "erp_sales": 100.0, "tbk_sales": 100.11, "difference": 0.11},
                {"card_code": "DI", "card_name": "Diners Club US$", "erp_sales": 0.0, "tbk_sales": 0.0, "difference": 0.0},
            ],
            "clp_table": [
                {"card_code": "VI", "card_name": "Visa (CLP)", "erp_sales": 95000.0, "tbk_sales": 95000.0, "difference": 0.0},
                {"card_code": "DI", "card_name": "Diners Club (CLP)", "erp_sales": 0.0, "tbk_sales": 0.0, "difference": 0.0},
            ],
            "usd_totals": {"erp": 100.0, "tbk": 100.11, "diff": 0.11},
            "clp_totals": {"erp": 95000.0, "tbk": 95000.0, "diff": 0.0},
            "matched_vouchers": [voucher],
            "discrepancy_vouchers": [],
            "orphan_vouchers": [],
            "counts": {"matched": 1, "discrepancies": 0, "orphans": 0, "total": 1},
        })

        response = self.client.get(reverse("usd_preview"))

        self.assertEqual(response.status_code, 200)
        rows = {row["card_code"]: row for row in response.context["franchise_rows"]}
        self.assertEqual(rows["VI"]["card_name"], "Visa")
        self.assertEqual(rows["VI"]["clp"]["erp_sales"], 95000.0)
        self.assertTrue(rows["VI"]["has_activity"])
        self.assertFalse(rows["DI"]["has_activity"])
        self.assertEqual(response.context["match_rate"], 100.0)
        self.assertContains(response, "A10001")



class BruteForceProtectionTests(TestCase):
    """django-axes locks the username+IP pair after 5 failed logins (both /login/ and the admin)."""

    def setUp(self):
        self.user = User.objects.create_superuser("contadora", password="clave-correcta-123")

    def _login(self, url, password, ip="10.0.0.1"):
        return self.client.post(url, {"username": "contadora", "password": password}, HTTP_X_REAL_IP=ip)

    def test_correct_password_logs_in(self):
        self.assertRedirects(self._login(reverse("login"), "clave-correcta-123"), reverse("dashboard"))

    def test_lockout_after_failed_attempts_blocks_even_the_correct_password(self):
        for _ in range(4):
            self.assertEqual(self._login(reverse("login"), "mala").status_code, 200)
        # The 5th failure reaches AXES_FAILURE_LIMIT and is already answered with the lockout page
        self.assertEqual(self._login(reverse("login"), "mala").status_code, 429)

        response = self._login(reverse("login"), "clave-correcta-123")

        self.assertEqual(response.status_code, 429)
        self.assertContains(response, "Acceso bloqueado", status_code=429)
        self.assertNotIn("_auth_user_id", self.client.session)

    def test_failed_login_shows_remaining_attempts(self):
        self.assertContains(self._login(reverse("login"), "mala"), "quedan <strong>4 intentos</strong>", html=False)
        self._login(reverse("login"), "mala")
        self._login(reverse("login"), "mala")
        self.assertContains(self._login(reverse("login"), "mala"), "último intento")

    def test_lockout_is_per_ip_so_attacker_cannot_block_real_user(self):
        for _ in range(5):
            self._login(reverse("login"), "mala", ip="203.0.113.9")

        response = self._login(reverse("login"), "clave-correcta-123", ip="10.0.0.1")

        self.assertRedirects(response, reverse("dashboard"))

    def test_admin_login_is_also_protected(self):
        admin_login = reverse("admin:login")
        for _ in range(5):
            self._login(admin_login, "mala")

        self.assertEqual(self._login(admin_login, "clave-correcta-123").status_code, 429)


# ------------------------------------------------------------------------------
# Multi-File Reconciliation & Deduplication Tests
# ------------------------------------------------------------------------------

from .services.file_merger import (
    ensure_file_list,
    merge_and_parse_clp_bordero,
    merge_and_parse_clp_tbk,
    merge_and_parse_bank_statements,
    validate_multi_deposit_bank_movements,
)


class MultiFileReconciliationTests(TestCase):
    """
    Tests the multi-file merging and deduplication capabilities across extended date ranges.
    """

    def setUp(self):
        self.user = User.objects.create_user("auditor", password="password123")

    def test_ensure_file_list_handles_all_shapes(self):
        self.assertEqual(ensure_file_list(None), [])
        f1 = io.BytesIO(b"dummy1")
        self.assertEqual(ensure_file_list(f1), [f1])
        f2 = io.BytesIO(b"dummy2")
        self.assertEqual(ensure_file_list([f1, None, f2]), [f1, f2])

    def test_merge_and_parse_clp_tbk_aggregates_and_deduplicates(self):
        # File 1: 2 sales (voucher 1001, voucher 1002)
        sales_1 = [
            ("07/09/2026", "RESTAURANT", 25000, "VN", 0, 25000, 500, 95, "**** 1234", "1001"),
            ("07/09/2026", "BAR", 15000, "VN", 0, 15000, 300, 57, "**** 5678", "1002"),
        ]
        f1 = build_tbk_settlement("tbk_day1.xlsx", "Crédito", sales_1)

        # File 2: 2 sales (voucher 1002 duplicated from day 1, and voucher 1003 new)
        sales_2 = [
            ("07/09/2026", "BAR", 15000, "VN", 0, 15000, 300, 57, "**** 5678", "1002"),
            ("08/09/2026", "RECEPCION", 50000, "VN", 0, 50000, 1000, 190, "**** 9999", "1003"),
        ]
        f2 = build_tbk_settlement("tbk_day2.xlsx", "Crédito", sales_2)

        merged = merge_and_parse_clp_tbk({"CREDIT": [f1, f2]})
        self.assertIn("CREDIT", merged)
        credit = merged["CREDIT"]

        # Sales vouchers should be deduplicated: 1001, 1002, 1003 (3 vouchers total, not 4)
        sales = credit["sales"]
        self.assertEqual(len(sales), 3)
        auth_codes = [s["authorization_code"] for s in sales]
        self.assertIn("1001", auth_codes)
        self.assertIn("1002", auth_codes)
        self.assertIn("1003", auth_codes)

        # Header deposits should aggregate both settlement files
        header = credit["header"]
        self.assertEqual(header["file_count"], 2)
        self.assertGreater(header["total_deposit"], Decimal("0"))

    def test_merge_and_parse_clp_bordero_deduplicates_overlapping_rows(self):
        # Payment 1 and Payment 2 in file 1
        p1 = (101, 10001, "CNBQ", "Alojamiento", "F-101", 25000, "1001", "07/09/2026", "10:00", "cajero1", "Huesped A")
        p2 = (102, 10002, "CNBQ", "Bar", "F-102", 15000, "1002", "07/09/2026", "11:00", "cajero1", "Huesped B")
        f1 = build_bordero([p1, p2])

        # File 2 overlaps p2, and adds p3
        p3 = (103, 10003, "CNBQ", "Restaurant", "F-103", 50000, "1003", "08/09/2026", "12:00", "cajero2", "Huesped C")
        f2 = build_bordero([p2, p3])

        all_rows, rows_by_auth, (min_date, max_date) = merge_and_parse_clp_bordero([f1, f2])
        # 3 unique rows: p1, p2, p3
        self.assertEqual(len(all_rows), 3)
        self.assertIn("1001", rows_by_auth)
        self.assertIn("1002", rows_by_auth)
        self.assertIn("1003", rows_by_auth)

    def test_validate_multi_deposit_bank_movements_individual_matching(self):
        m1 = BankStatementMovement(movement_date=date(2026, 9, 8), document_number="DOC1", deposit_amount=Decimal("39100.00"), description="ABONO TBK")
        m2 = BankStatementMovement(movement_date=date(2026, 9, 8), document_number="DOC2", deposit_amount=Decimal("48810.00"), description="ABONO TBK")

        expected = [
            ("CREDIT_DAY1", Decimal("39100.00")),
            ("CREDIT_DAY2", Decimal("48810.00")),
        ]
        summary, claimed = validate_multi_deposit_bank_movements([m1, m2], expected, total_consolidated=Decimal("87910.00"))

        self.assertEqual(len(claimed), 2)
        self.assertTrue(all(v["is_matched"] for v in summary))
        self.assertFalse(any(v.get("is_consolidated_match") for v in summary))

    def test_validate_multi_deposit_bank_movements_consolidated_fallback(self):
        # Bank made one combined lump-sum deposit for the weekend: 39,100 + 48,810 = 87,910
        m_consolidated = BankStatementMovement(
            movement_date=date(2026, 9, 8),
            document_number="BULK1",
            deposit_amount=Decimal("87910.00"),
            description="ABONO TRANSBANK FIN DE SEMANA",
        )

        expected = [
            ("CREDIT_DAY1", Decimal("39100.00")),
            ("CREDIT_DAY2", Decimal("48810.00")),
        ]
        summary, claimed = validate_multi_deposit_bank_movements(
            [m_consolidated],
            expected,
            total_consolidated=Decimal("87910.00"),
        )

        self.assertEqual(len(claimed), 1)
        self.assertTrue(all(v["is_matched"] for v in summary))
        self.assertTrue(all(v.get("is_consolidated_match") for v in summary))

    def test_clp_preview_accepts_file_lists(self):
        sales_credit = [("07/09/2026", "RESTAURANT", 25000, "VN", 0, 25000, 500, 95, "**** 1234", "A-101")]
        sales_debit = [("07/09/2026", "BAR", 15000, "VD", 0, 15000, 300, 57, "**** 5678", "A-102")]

        f_tbk_c = build_tbk_settlement("tbk_c.xlsx", "Crédito", sales_credit)
        f_tbk_d = build_tbk_settlement("tbk_d.xlsx", "Redcompra Débito", sales_debit)
        f_tbk_p = build_tbk_settlement("tbk_p.xlsx", "Prepago", [])

        p1 = (101, 10001, "CNBQ", "Alojamiento", "F-101", 25000, "A-101", "07/09/2026", "10:00", "cajero1", "Huesped A")
        p2 = (102, 10002, "VN", "Bar", "F-102", 15000, "A-102", "07/09/2026", "11:00", "cajero1", "Huesped B")
        f_bordero1 = build_bordero([p1])
        f_bordero2 = build_bordero([p2])

        cartola_dep = [
            ("08/09/2026", "TRANSBANK CREDITO", 24405),
            ("08/09/2026", "TRANSBANK DEBITO", 14643),
        ]
        f_cartola = build_cartola(cartola_dep)

        preview = calculate_clp_preview(
            erp_files=[f_bordero1, f_bordero2],
            bank_statement_files=[f_cartola],
            tbk_credit_files=[f_tbk_c],
            tbk_debit_files=[f_tbk_d],
            tbk_prepaid_files=[f_tbk_p],
        )

        self.assertIn("summary", preview)
        self.assertIn("matched_vouchers", preview)
        self.assertEqual(preview["counts"]["matched"], 2)
        self.assertEqual(preview["counts"]["discrepancies"], 0)
